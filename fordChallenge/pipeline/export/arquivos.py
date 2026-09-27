"""Os arquivos: XLSX, Parquet e o DDL do BigQuery.

**Cada formato aqui tem um destinatário diferente, e é isso que justifica os três.** O CSV
(em `tabela.py`) é para quem vai abrir e olhar. O XLSX é para quem vai *trabalhar* em cima
— largura de coluna, cabeçalho congelado, uma aba separada para as fontes. O Parquet e o
DDL são para quem vai **carregar na stack da Ford**: ali ninguém lê a tabela, um pipeline
a ingere.

**Dependência que falta não vira arquivo vazio nem exceção crua.** Vira
:class:`FormatoIndisponivel`, com o nome do extra a instalar. `openpyxl` e `pyarrow` não
estão na instalação base de propósito: são pesados, e o caminho crítico da apresentação
não pode depender deles. Quem exporta num ambiente sem eles recebe a frase certa, e não
uma planilha de zero byte que só se descobre quebrada na frente de outra pessoa.
"""

from __future__ import annotations

from pathlib import Path

from pipeline.export.tabela import Tabela


class FormatoIndisponivel(RuntimeError):
    """O formato existe, a biblioteca não está instalada. Diz qual e como instalar."""


def _exigir(modulo: str, extra: str):
    try:
        return __import__(modulo)
    except ModuleNotFoundError as exc:  # pragma: no cover - depende do ambiente
        raise FormatoIndisponivel(
            f"{modulo} não está instalado. Rode `uv sync --extra {extra}` e tente de novo. "
            "Os outros formatos continuam disponíveis."
        ) from exc


def para_xlsx(tabela: Tabela, destino: Path) -> Path:
    """A planilha: uma aba com a matriz, outra com as fontes e as ressalvas.

    **Duas abas, e não uma com tudo empilhado.** A matriz é para ler em linha; a lista de
    fontes é para conferir uma URL específica. Misturar as duas faz a planilha rolar para
    o lado e para baixo ao mesmo tempo, e é assim que uma exportação vira um anexo que
    ninguém abre duas vezes.
    """
    openpyxl = _exigir("openpyxl", "export")
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    livro = openpyxl.Workbook()
    aba = livro.active
    aba.title = "matriz"

    aba.append([tabela.titulo])
    aba["A1"].font = Font(bold=True, size=13)
    aba.append([f"gerado em {tabela.gerada_em}"])
    if tabela.versao_dos_criterios:
        aba.append([f"critérios de comparabilidade: versão {tabela.versao_dos_criterios}"])
    aba.append([])

    linha_do_cabecalho = aba.max_row + 1
    aba.append(tabela.cabecalho())
    for celula in aba[linha_do_cabecalho]:
        celula.font = Font(bold=True)
        celula.alignment = Alignment(vertical="center", wrap_text=True)

    for linha in tabela.como_linhas():
        aba.append(linha)

    # Congelar abaixo do cabeçalho **e** à direita da primeira coluna: sem isso, rolar
    # para a quarta coluna de veículo esconde o nome do campo, e a pessoa lê o número
    # sem saber de qual campo ele é — que é pior do que não ver o número.
    aba.freeze_panes = aba.cell(row=linha_do_cabecalho + 1, column=2)
    larguras = [46, 18] + [30, 16, 52, 14] * len(tabela.colunas)
    for indice, largura in enumerate(larguras[: aba.max_column], start=1):
        aba.column_dimensions[get_column_letter(indice)].width = largura

    fontes = livro.create_sheet("fontes e ressalvas")
    fontes.append(["veículo", "comparabilidade", "aviso"])
    for celula in fontes[1]:
        celula.font = Font(bold=True)
    for coluna in tabela.colunas:
        fontes.append([coluna.rotulo, coluna.comparabilidade or "—", coluna.aviso or "—"])
    fontes.append([])
    fontes.append(["ressalvas"])
    fontes.cell(row=fontes.max_row, column=1).font = Font(bold=True)
    for ressalva in tabela.ressalvas or ["nenhuma"]:
        fontes.append([ressalva])
    for indice, largura in enumerate((38, 26, 90), start=1):
        fontes.column_dimensions[get_column_letter(indice)].width = largura

    destino.parent.mkdir(parents=True, exist_ok=True)
    livro.save(destino)
    return destino


#: As colunas do Parquet. **Formato longo**, e não a matriz larga do CSV.
#:
#: A matriz larga tem uma coluna por veículo, e por isso o esquema **muda** a cada
#: comparação — três concorrentes hoje, cinco amanhã. Uma tabela cujo esquema muda não se
#: carrega em warehouse: cada exportação exigiria uma migração. No formato longo o esquema
#: é fixo e o veículo vira dado, que é como toda a stack espera receber.
COLUNAS_DO_PARQUET: tuple[str, ...] = (
    "version_id",
    "veiculo",
    "e_referencia",
    "campo",
    "grupo",
    "valor",
    "unidade",
    "status",
    "situacao",
    "fonte_url",
    "fonte_data",
    "trecho",
    "etiqueta",
    "gerado_em",
)


def linhas_longas(tabela: Tabela) -> list[dict[str, object]]:
    """A tabela em formato longo: uma linha por (veículo, campo). Esquema fixo."""
    saida: list[dict[str, object]] = []
    for linha in tabela.linhas:
        for coluna, celula in zip(tabela.colunas, linha.celulas, strict=False):
            saida.append(
                {
                    "version_id": coluna.version_id,
                    "veiculo": coluna.rotulo,
                    "e_referencia": coluna.e_referencia,
                    "campo": linha.campo,
                    "grupo": linha.grupo,
                    "valor": celula.texto,
                    "unidade": celula.unidade,
                    "status": celula.status,
                    "situacao": celula.situacao,
                    "fonte_url": celula.fonte_url,
                    "fonte_data": celula.fonte_data,
                    "trecho": celula.trecho,
                    "etiqueta": celula.etiqueta,
                    "gerado_em": tabela.gerada_em,
                }
            )
    return saida


def para_parquet(tabela: Tabela, destino: Path) -> Path:
    """O Parquet, em formato longo. Exige `pyarrow`."""
    _exigir("pyarrow", "export")
    import pandas as pd

    destino.parent.mkdir(parents=True, exist_ok=True)
    quadro = pd.DataFrame(linhas_longas(tabela), columns=list(COLUNAS_DO_PARQUET))
    quadro.to_parquet(destino, index=False)
    return destino


#: O tipo BigQuery de cada coluna. Tudo STRING menos o que **é** outra coisa.
#:
#: `valor` é STRING de propósito, e não FLOAT: a mesma coluna guarda 397, "sim", "2,8" e
#: "não encontrado em nenhuma fonte". Tipar como número obrigaria a jogar fora as três
#: últimas — e são elas que distinguem esta ficha de uma planilha qualquer.
TIPOS_DO_BIGQUERY: dict[str, str] = {
    "version_id": "STRING",
    "veiculo": "STRING",
    "e_referencia": "BOOL",
    "campo": "STRING",
    "grupo": "STRING",
    "valor": "STRING",
    "unidade": "STRING",
    "status": "STRING",
    "situacao": "STRING",
    "fonte_url": "STRING",
    "fonte_data": "STRING",
    "trecho": "STRING",
    "etiqueta": "STRING",
    "gerado_em": "TIMESTAMP",
}


def ddl_do_bigquery(*, dataset: str = "specradar", tabela_nome: str = "matriz_de_paridade") -> str:
    """O `CREATE TABLE` do BigQuery para o esquema longo.

    Particionado por `gerado_em` e agrupado por `campo`: as duas consultas que essa tabela
    recebe são "como estava em tal data" e "a evolução deste campo". Sem a partição, cada
    consulta varre o histórico inteiro e a conta de quem hospeda cresce por descuido nosso.
    """
    colunas = ",\n".join(
        f"  {nome} {tipo}" + (" NOT NULL" if nome in {"version_id", "campo"} else "")
        for nome, tipo in TIPOS_DO_BIGQUERY.items()
    )
    return (
        f"CREATE TABLE IF NOT EXISTS `{dataset}.{tabela_nome}` (\n"
        f"{colunas}\n"
        ")\n"
        "PARTITION BY DATE(gerado_em)\n"
        "CLUSTER BY campo\n"
        "OPTIONS (\n"
        '  description = "SpecRadar — matriz de paridade em formato longo. '
        'Uma linha por (veiculo, campo), com a fonte e a data da captura."\n'
        ");\n"
    )
